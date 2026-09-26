"""
Document Ingestion & Parsing Engine.
Validates uploaded files, extracts structured text from multi-page PDFs and text files,
detects scanned/image-only documents requiring OCR, and creates deterministic text chunks
with rich metadata for vector indexing.
"""
import os
import re
import uuid
import hashlib
import logging
from pathlib import Path
from typing import List, Dict, Any, Optional
import pypdf

from app.config import settings

logger = logging.getLogger("vaxassist.knowledge.parser")


class EmptyDocumentError(ValueError):
    """Raised when a document contains zero or negligible extractable text."""
    pass


class DocumentParser:
    """
    Handles file validation, text extraction, normalization, and chunk generation.
    """

    @staticmethod
    def calculate_sha256(content: bytes) -> str:
        """Computes SHA-256 hash of file bytes for integrity tracking."""
        return hashlib.sha256(content).hexdigest()

    @staticmethod
    def validate_and_save_file(
        filename: str,
        content: bytes,
        target_dir: Optional[Path] = None,
    ) -> Dict[str, Any]:
        """
        Validates uploaded file size and extension, computes safe filename, and persists to disk.
        Prevents path traversal and unsafe filenames.
        """
        dest_dir = target_dir or Path(settings.KNOWLEDGE_STORAGE_DIRECTORY)
        dest_dir.mkdir(parents=True, exist_ok=True)

        # 1. Size validation
        size_bytes = len(content)
        if size_bytes == 0:
            raise ValueError("Uploaded file is empty (0 bytes).")
        if size_bytes > settings.MAX_UPLOAD_FILE_SIZE_BYTES:
            max_mb = settings.MAX_UPLOAD_FILE_SIZE_BYTES / (1024 * 1024)
            raise ValueError(f"File size exceeds maximum permitted limit of {max_mb:.0f} MB.")

        # 2. Extension validation
        orig_name = Path(filename).name  # Prevent directory path traversal in filename
        ext = Path(orig_name).suffix.lower()
        if ext not in settings.ALLOWED_DOCUMENT_EXTENSIONS:
            allowed = ", ".join(settings.ALLOWED_DOCUMENT_EXTENSIONS)
            raise ValueError(f"Unsupported file format '{ext}'. Allowed formats: {allowed}")

        # 3. Determine MIME type
        if ext == ".pdf":
            mime_type = "application/pdf"
        elif ext == ".docx":
            mime_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        else:
            mime_type = "text/plain"

        # 4. Generate collision-free, safe stored filename
        unique_id = uuid.uuid4().hex
        stored_filename = f"{unique_id}{ext}"
        saved_path = dest_dir / stored_filename

        # 5. Persist file bytes
        with open(saved_path, "wb") as f:
            f.write(content)

        file_hash = DocumentParser.calculate_sha256(content)

        logger.info(f"Persisted document '{orig_name}' -> '{saved_path}' ({size_bytes} bytes, hash: {file_hash[:8]}...)")

        return {
            "original_filename": orig_name,
            "stored_filename": stored_filename,
            "file_path": str(saved_path),
            "file_size_bytes": size_bytes,
            "file_hash": file_hash,
            "mime_type": mime_type,
        }

    @staticmethod
    def extract_pages(file_path: str) -> List[Dict[str, Any]]:
        """
        Extracts clean text page-by-page from a PDF or text file.
        Detects scanned or image-only PDFs that contain no text layer.
        """
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"Document file not found at: {file_path}")

        ext = path.suffix.lower()
        pages: List[Dict[str, Any]] = []

        if ext == ".pdf":
            try:
                reader = pypdf.PdfReader(str(path))
                if reader.is_encrypted:
                    try:
                        reader.decrypt("")
                    except Exception:
                        raise ValueError("Encrypted PDF documents are not supported. Please upload an unencrypted copy.")

                total_chars = 0
                for idx, page in enumerate(reader.pages):
                    raw_text = page.extract_text() or ""
                    cleaned = DocumentParser._normalize_text(raw_text)
                    total_chars += len(cleaned)
                    if cleaned:
                        pages.append({
                            "page_number": idx + 1,
                            "text": cleaned,
                        })

                # Scanned / Image-only verification
                if total_chars < 50:
                    raise EmptyDocumentError(
                        "Scanned or image-only PDF detected. Text extraction requires OCR which is not supported "
                        "in this version. Please upload text-based PDF or document."
                    )

            except EmptyDocumentError:
                raise
            except Exception as e:
                logger.error(f"Error reading PDF {file_path}: {e}")
                raise ValueError(f"Failed to extract text from PDF document: {str(e)}")

        elif ext == ".docx":
            try:
                import docx
                doc = docx.Document(str(path))
                full_text = []
                for p in doc.paragraphs:
                    if p.text.strip():
                        full_text.append(p.text.strip())
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            full_text.append(row_text)
                cleaned = DocumentParser._normalize_text("\n\n".join(full_text))
                if len(cleaned) < 20:
                    raise EmptyDocumentError("DOCX document contains insufficient or empty text.")
                pages.append({
                    "page_number": 1,
                    "text": cleaned,
                })
            except EmptyDocumentError:
                raise
            except Exception as e:
                logger.error(f"Error reading DOCX {file_path}: {e}")
                raise ValueError(f"Failed to extract text from DOCX document: {str(e)}")

        else:
            # Plain text or Markdown
            try:
                content = path.read_text(encoding="utf-8")
            except UnicodeDecodeError:
                content = path.read_text(encoding="latin-1")

            cleaned = DocumentParser._normalize_text(content)
            if len(cleaned) < 20:
                raise EmptyDocumentError("Text document contains insufficient or empty text.")

            pages.append({
                "page_number": 1,
                "text": cleaned,
            })

        return pages

    @staticmethod
    def _normalize_text(text: str) -> str:
        """Cleans whitespace, line breaks, and non-printable characters."""
        if not text:
            return ""
        # Remove null bytes
        text = text.replace("\x00", "")
        # Normalize newlines
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        # Collapse excessive newlines and spaces
        text = re.sub(r"\n{3,}", "\n\n", text)
        text = re.sub(r"[ \t]+", " ", text)
        return text.strip()

    @staticmethod
    def chunk_document(
        pages: List[Dict[str, Any]],
        document_id: str,
        document_title: str,
        source_authority: str,
        source_url: Optional[str] = None,
        index_version: int = 1,
        chunk_size: Optional[int] = None,
        chunk_overlap: Optional[int] = None,
    ) -> List[Dict[str, Any]]:
        """
        Generates deterministic text chunks preserving page numbers and document context.
        Splits on sentence and paragraph boundaries to preserve clinical and semantic coherence.
        """
        size = chunk_size or settings.CHUNK_SIZE_CHARS
        overlap = chunk_overlap or settings.CHUNK_OVERLAP_CHARS

        chunks: List[Dict[str, Any]] = []
        global_chunk_idx = 0

        for page in pages:
            page_num = page["page_number"]
            page_text = page["text"]

            if len(page_text) <= size:
                # Small page fits in single chunk
                chunk_id = f"{document_id}_v{index_version}_c{global_chunk_idx}"
                chunks.append({
                    "id": chunk_id,
                    "content": page_text,
                    "metadata": {
                        "knowledge_document_id": str(document_id),
                        "document_title": document_title,
                        "source_authority": source_authority,
                        "source_url": source_url or "",
                        "page_number": page_num,
                        "chunk_index": global_chunk_idx,
                        "index_version": index_version,
                    }
                })
                global_chunk_idx += 1
                continue

            # Split larger text into overlapping window segments
            start = 0
            text_len = len(page_text)
            while start < text_len:
                end = min(start + size, text_len)

                # Try to break at paragraph boundary, sentence end, or whitespace
                if end < text_len:
                    break_point = -1
                    # Check paragraph boundary
                    p_break = page_text.rfind("\n\n", start + int(size * 0.5), end)
                    if p_break != -1:
                        break_point = p_break + 2
                    else:
                        # Check sentence end
                        s_break = page_text.rfind(". ", start + int(size * 0.5), end)
                        if s_break != -1:
                            break_point = s_break + 2
                        else:
                            # Check whitespace
                            w_break = page_text.rfind(" ", start + int(size * 0.7), end)
                            if w_break != -1:
                                break_point = w_break + 1

                    if break_point > start:
                        end = break_point

                chunk_content = page_text[start:end].strip()
                if len(chunk_content) >= 30:  # Skip trivially tiny leftover fragments
                    chunk_id = f"{document_id}_v{index_version}_c{global_chunk_idx}"
                    chunks.append({
                        "id": chunk_id,
                        "content": chunk_content,
                        "metadata": {
                            "knowledge_document_id": str(document_id),
                            "document_title": document_title,
                            "source_authority": source_authority,
                            "source_url": source_url or "",
                            "page_number": page_num,
                            "chunk_index": global_chunk_idx,
                            "index_version": index_version,
                        }
                    })
                    global_chunk_idx += 1

                start = max(start + 1, end - overlap)

        return chunks


document_parser = DocumentParser()
